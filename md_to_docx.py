import sys
import docx

def create_docx(md_text, output_path):
    doc = docx.Document()
    
    lines = md_text.split('\n')
    for line in lines:
        if line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('**') and line.endswith('**'):
            p = doc.add_paragraph()
            p.add_run(line[2:-2]).bold = True
        elif line.strip() == '---':
            pass
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif line.strip() == '':
            pass
        else:
            doc.add_paragraph(line)

    doc.save(output_path)
    print(f"Document saved to {output_path}")

if __name__ == '__main__':
    md_file = sys.argv[1]
    out_file = sys.argv[2]
    
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
        
    create_docx(content, out_file)
